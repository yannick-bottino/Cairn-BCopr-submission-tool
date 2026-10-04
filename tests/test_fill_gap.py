import json
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills" / "bcorp-gap-analysis" / "scripts"))

import build_gap_excel as bg  # noqa: E402
import fill_gap as fg  # noqa: E402
from test_build_gap_excel import PROFILE  # noqa: E402

DIAG = "Diagnostic & gap analysis"
COMMENT = "Commentaire soumission dossier pour l'auditeur"


@pytest.fixture
def gap_file(tmp_path):
    return bg.build(PROFILE, tmp_path / "gap.xlsx")


def rows_of(path):
    ws = openpyxl.load_workbook(path)["3. Gap analysis"]
    headers = [c.value for c in ws[1]]
    return {r[headers.index("req_id")].value: dict(zip(headers, [c.value for c in r]))
            for r in ws.iter_rows(min_row=2)}


def proposal(req_id="PSG1.1-1.1.1", **fields):
    return {"req_id": req_id, "champs": fields or {DIAG: "Statuts du 15/01/2024, art. 2 : raison d'être inscrite."}}


def test_writes_into_new_file_only(gap_file, tmp_path):
    before = gap_file.read_bytes()
    out = tmp_path / "gap_v2.xlsx"
    report = fg.fill(gap_file, [proposal()], out)
    assert gap_file.read_bytes() == before
    assert rows_of(out)["PSG1.1-1.1.1"][DIAG].startswith("Statuts du 15/01/2024")
    assert report.written == 1


def test_refuses_to_overwrite_output(gap_file):
    with pytest.raises(FileExistsError):
        fg.fill(gap_file, [proposal()], gap_file)


def test_never_overwrites_a_filled_cell(gap_file, tmp_path):
    first = fg.fill(gap_file, [proposal()], tmp_path / "v2.xlsx")
    assert first.written == 1
    second = fg.fill(tmp_path / "v2.xlsx", [proposal(**{DIAG: "Autre texte"})], tmp_path / "v3.xlsx")
    assert second.written == 0 and len(second.skipped) == 1
    assert rows_of(tmp_path / "v3.xlsx")["PSG1.1-1.1.1"][DIAG].startswith("Statuts du 15/01/2024")


def test_only_writer_columns(gap_file, tmp_path):
    for col in ("Responsable", "Deadline", "Statut", "Plateforme : statut", "Commentaires", "Code exigence"):
        with pytest.raises(ValueError, match="non modifiable"):
            fg.fill(gap_file, [proposal(**{col: "x"})], tmp_path / f"{col[:5]}.xlsx")


def test_unknown_req_id(gap_file, tmp_path):
    with pytest.raises(KeyError):
        fg.fill(gap_file, [proposal(req_id="PSG9.9-9.9.9")], tmp_path / "x.xlsx")


def test_list_values_are_checked(gap_file, tmp_path):
    with pytest.raises(ValueError, match="Niveau de conformité"):
        fg.fill(gap_file, [proposal(**{"Niveau de conformité": "3"})], tmp_path / "x.xlsx")
    ok = fg.fill(gap_file, [proposal(**{"Niveau de conformité": "1", "Criticité (Anchor)": "2-Majeur"})],
                 tmp_path / "y.xlsx")
    assert ok.written == 2


def test_invented_code_in_text_is_refused(gap_file, tmp_path):
    with pytest.raises(ValueError, match="CA4.4"):
        fg.fill(gap_file, [proposal(**{DIAG: "Voir aussi CA4.4 pour le plan climat."})], tmp_path / "x.xlsx")
    ok = fg.fill(gap_file, [proposal(**{DIAG: "Voie la plus directe : APAC 1.1.1, comme pour PSG1.1."})],
                 tmp_path / "y.xlsx")
    assert ok.written == 1


@pytest.mark.parametrize("text, rule", [
    ("La raison d'être — inscrite aux statuts.", "tiret cadratin"),
    ("**Conforme** : statuts.", "Markdown"),
])
def test_style_errors_block(gap_file, tmp_path, text, rule):
    with pytest.raises(ValueError, match=rule):
        fg.fill(gap_file, [proposal(**{DIAG: text})], tmp_path / "x.xlsx")


def test_auditor_comment_never_states_a_gap(gap_file, tmp_path):
    for bad in ("Manquant : la politique.", "Question à poser : qui valide ?"):
        with pytest.raises(ValueError, match="commentaire auditeur"):
            fg.fill(gap_file, [proposal(**{COMMENT: bad})], tmp_path / "x.xlsx")


def test_in_progress_without_milestone_is_a_warning(gap_file, tmp_path):
    report = fg.fill(gap_file, [proposal(**{COMMENT: "Le dispositif est en cours de déploiement.\n- Pièces jointes : charte"})],
                     tmp_path / "x.xlsx")
    assert report.written == 1
    assert any("jalon" in w for w in report.warnings)


def test_proposals_are_logged_for_validation(gap_file, tmp_path):
    out = tmp_path / "v2.xlsx"
    fg.fill(gap_file, [proposal()], out)
    wb = openpyxl.load_workbook(out)
    ws = wb["Propositions"]
    headers = [c.value for c in ws[1]]
    assert headers[:4] == ["req_id", "Code plateforme", "Colonne", "Statut"]
    row = [c.value for c in ws[2]]
    assert row[0] == "PSG1.1-1.1.1" and row[2] == DIAG and row[3] == "À valider"
    gap = wb["3. Gap analysis"]
    col = [c.value for c in gap[1]].index(DIAG) + 1
    cell = next(gap.cell(r, col) for r in range(2, gap.max_row + 1) if gap.cell(r, col).value)
    assert cell.fill.fgColor.rgb.endswith(fg.PROPOSAL_COLOR)


def test_cli_reads_json(gap_file, tmp_path):
    p = tmp_path / "props.json"
    p.write_text(json.dumps([proposal()], ensure_ascii=False), encoding="utf-8")
    assert fg.main([str(gap_file), str(p), str(tmp_path / "cli.xlsx")]) == 0


@pytest.mark.parametrize("text", ["Option retenue : JEDI 2.g.", "Voir GACA2.3c et ESC3.3a.", "Critère JEDI2.a.1."])
def test_option_codes_accepted(gap_file, tmp_path, text):
    assert fg.fill(gap_file, [proposal(**{DIAG: text})], tmp_path / "x.xlsx").written == 1


@pytest.mark.parametrize("text", ["Option JEDI 2.zz.", "Voir GACA2.3z.", "Critère PSG1.1.9."])
def test_unknown_option_or_criterion_refused(gap_file, tmp_path, text):
    with pytest.raises(ValueError, match="Code inconnu"):
        fg.fill(gap_file, [proposal(**{DIAG: text})], tmp_path / "x.xlsx")


def test_revenue_figures_are_not_codes(gap_file, tmp_path):
    ok = fg.fill(gap_file, [proposal(**{DIAG: "CA 2024 : 33,5 M€ ; FR 2 sites."})], tmp_path / "x.xlsx")
    assert ok.written == 1
