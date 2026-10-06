import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills" / "bcorp-dataroom-audit" / "scripts"))

import inventory as inv  # noqa: E402


@pytest.fixture
def dataroom(tmp_path):
    d = tmp_path / "Inputs client"
    for rel, content in {
        "Juridique/Statuts ClientTest 2024.pdf": b"statuts",
        "Juridique/Kbis.pdf": b"kbis",
        "RSE/Bilan carbone 2024.xlsx": b"bc",
        "RSE/Bilan carbone 2024 copie.xlsx": b"bc",                 # doublon exact
        "RH/Charte télétravail v1.docx": b"v1",
        "RH/Charte télétravail v2.docx": b"v2",                    # v1 obsolète
        "Divers/photo.jpg": b"img",
        ".DS_Store": b"x",
    }.items():
        p = d / rel
        p.parent.mkdir(parents=True, exist_ok=True)
        p.write_bytes(content)
    return d


def sheet(path, name):
    ws = openpyxl.load_workbook(path)[name]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


def test_inventory_lists_every_file_read_only(dataroom, tmp_path):
    before = sorted(p.relative_to(dataroom) for p in dataroom.rglob("*"))
    out = inv.build(dataroom, tmp_path / "inventaire.xlsx")
    assert sorted(p.relative_to(dataroom) for p in dataroom.rglob("*")) == before
    rows = sheet(out, "Inventaire")
    assert len(rows) == 7 and not any(".DS_Store" in r["Chemin"] for r in rows)


def test_classification_is_a_pre_sort(dataroom, tmp_path):
    rows = {r["Fichier"]: r for r in sheet(inv.build(dataroom, tmp_path / "i.xlsx"), "Inventaire")}
    assert rows["Statuts ClientTest 2024.pdf"]["Type pressenti"] == "Juridique"
    assert "FR" in rows["Statuts ClientTest 2024.pdf"]["Impact Areas pressenties"]
    assert "CA" in rows["Bilan carbone 2024.xlsx"]["Impact Areas pressenties"]
    assert rows["photo.jpg"]["Type pressenti"] == "À classer"


def test_duplicates_and_obsolete_versions(dataroom, tmp_path):
    rows = {r["Fichier"]: r for r in sheet(inv.build(dataroom, tmp_path / "i.xlsx"), "Inventaire")}
    assert "doublon" in (rows["Bilan carbone 2024 copie.xlsx"]["Alerte"] or "")
    assert "version antérieure" in (rows["Charte télétravail v1.docx"]["Alerte"] or "")
    assert not rows["Charte télétravail v2.docx"]["Alerte"]


def test_key_documents_checklist(dataroom, tmp_path):
    rows = {r["Pièce"]: r for r in sheet(inv.build(dataroom, tmp_path / "i.xlsx"), "Pièces clés")}
    assert rows["Extrait Kbis récent"]["Statut"] == "Présent (à vérifier)"
    assert rows["Table de capitalisation / actionnariat"]["Statut"] == "Absent"
    assert rows["Liasses fiscales ou comptes annuels"]["Statut"] == "Absent"


def test_pillar_view_has_empty_columns_for_claude(dataroom, tmp_path):
    rows = sheet(inv.build(dataroom, tmp_path / "i.xlsx"), "Vue par pilier")
    assert [r["Impact Area"].split(" (")[0] for r in rows][:2] == ["Exigences de base", "Mission et gouvernance des parties prenantes"]
    assert len(rows) == 8
    assert all(r["Ce qui existe (preuve)"] is None and r["Ce qui manque (Year 0)"] is None for r in rows)
    fr = rows[0]
    assert fr["Documents pressentis"] >= 2


def test_never_writes_inside_client_folder(dataroom, tmp_path):
    with pytest.raises(ValueError, match="dossier client"):
        inv.build(dataroom, dataroom / "inventaire.xlsx")
    out = tmp_path / "i.xlsx"
    out.write_bytes(b"x")
    with pytest.raises(FileExistsError):
        inv.build(dataroom, out)
