import sys
from pathlib import Path

import openpyxl
import pytest

import bcorp_ref as ref

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills" / "bcorp-gap-analysis" / "scripts"))
import build_gap_excel as bg  # noqa: E402
from test_build_gap_excel import PROFILE  # noqa: E402


@pytest.fixture(scope="module")
def r():
    return ref.Referentiel.load()


def test_every_risk_criterion_is_understood(r):
    rules = r.risk_rules()
    assert len(rules) == sum(len(x.criteria) for x in r.rows if x.type_ligne == "question_risk_tool")
    for rule in rules:
        assert rule.sizes and rule.sectors, rule.criterion
        assert rule.adds or rule.replaces or rule.consider or rule.no_change, rule.criterion
        for code in rule.adds + list(rule.replaces) + list(rule.replaces.values()) + rule.consider:
            r.to_site(code)                        # aucun code inventé


@pytest.mark.parametrize("question, size, sector, adds, replaces", [
    ("FR3.1.a", "Small", "Wholesale/Retail", {"HR2.1", "HR2.3", "HR2.4"}, {}),
    ("FR3.1.a", "Medium", "Wholesale/Retail", {"HR2.5"}, {}),
    ("FR3.1.a", "Large", "Service with Minor Environmental Footprint", {"HR2.3", "HR2.4", "HR2.5"}, {}),
    ("FR3.1.n", "Small", "Manufacturing", {"PSG4.2"}, {"PSG4.1": "PSG4.2"}),
    ("FR3.1.n", "Medium", "Wholesale/Retail", {"PSG3.3", "PSG3.4"}, {"PSG3.1": "PSG3.3", "PSG3.2": "PSG3.4"}),
    ("FR3.1.n", "Large", "Wholesale/Retail", set(), {}),
    ("FR3.1.l", "Medium", "Service with Minor Environmental Footprint", {"ESC1.4", "ESC1.7", "ESC2.1", "ESC4.2"}, {}),
    ("FR3.1.e", "Large", "Manufacturing", {"ESC3.4", "ESC3.5"}, {}),
])
def test_effects(r, question, size, sector, adds, replaces):
    eff = r.risk_effects(size, sector, "", [question])
    assert eff.adds == adds
    assert eff.replaces == replaces


def test_mining_industry(r):
    mining = r.risk_effects("Medium", "Manufacturing", "Mining", ["FR3.1.f"])
    other = r.risk_effects("Medium", "Manufacturing", "", ["FR3.1.f"])
    assert mining.adds == set() and "ESC1.7" in mining.consider
    assert other.adds == {"ESC2.2", "ESC4.3"}


def test_potential_impact_note(r):
    eff = r.risk_effects("XX Large", "Wholesale/Retail", "", ["FR3.1.a", "FR3.1.m"])
    assert eff.adds == set()
    assert eff.consider["HR2.1"] == ["FR3.1.a", "FR3.1.m"]


def test_micro_has_no_effect(r):
    assert r.risk_effects("Micro", "Wholesale/Retail", "", ["FR3.1.a"]).adds == set()


# --- intégration dans l'Excel --------------------------------------------------

def gap_rows(path):
    ws = openpyxl.load_workbook(path)["3. Gap analysis"]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, [c.value for c in row])) for row in ws.iter_rows(min_row=2)]


def test_gap_adds_and_replaces(tmp_path, r):
    base = {row["Code plateforme"] for row in gap_rows(bg.build(PROFILE, tmp_path / "a.xlsx"))}
    assert "PSG3.3" not in base
    rows = gap_rows(bg.build({**PROFILE, "risk_tool_oui": ["FR3.1.n"]}, tmp_path / "b.xlsx"))
    by_code = {}
    for row in rows:
        by_code.setdefault(row["Code plateforme"], []).append(row)
    assert "PSG3.3" in by_code and "PSG3.4" in by_code
    assert "Risk Tool" in by_code["PSG3.3"][0]["Clarification / informations complémentaires"]
    replaced = by_code["PSG3.1"][0]
    assert replaced["Niveau de conformité"] == "NA"
    assert "PSG3.3" in replaced["Clarification / informations complémentaires"]


def test_gap_marks_potential_impact(tmp_path):
    rows = gap_rows(bg.build({**PROFILE, "risk_tool_oui": ["FR3.1.c"]}, tmp_path / "c.xlsx"))
    hr21 = [row for row in rows if row["Code plateforme"] == "HR2.1"]
    assert hr21 and all("FR3.1.c" in (row["Clarification / informations complémentaires"] or "") for row in hr21)


def test_unknown_risk_question_refused(tmp_path):
    with pytest.raises(ref.UnknownCode):
        bg.build({**PROFILE, "risk_tool_oui": ["FR3.1.z"]}, tmp_path / "d.xlsx")
