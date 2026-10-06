import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
for d in ("bcorp-gap-analysis", "bcorp-evidence-pack", "bcorp-platform-plan"):
    sys.path.insert(0, str(ROOT / "skills" / d / "scripts"))

import build_gap_excel as bg  # noqa: E402
import fill_gap as fg  # noqa: E402
import evidence_plan as ep  # noqa: E402
import platform_plan as pp  # noqa: E402
from test_build_gap_excel import PROFILE  # noqa: E402

COMMENT = fg.COMMENT


def set_status(path, sheet, match, status, col="Statut"):
    wb = openpyxl.load_workbook(path)
    ws = wb[sheet]
    headers = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        values = {h: ws.cell(r, i + 1).value for i, h in enumerate(headers)}
        if match(values):
            ws.cell(r, headers.index(col) + 1).value = status
    wb.save(path)


@pytest.fixture
def gap(tmp_path):
    src = bg.build(PROFILE, tmp_path / "gap.xlsx")
    out = tmp_path / "gap_rempli.xlsx"
    fg.fill(src, [
        {"req_id": "PSG1.1-1.1.1", "champs": {COMMENT: "La raison d'être couvre les 4 critères a/b/c/d :\n- Statuts art. 2\n- Pièces jointes : statuts", "Preuves (& intitulés associés)": "- Statuts ClientTest"}},
        {"req_id": "PSG2.1-2.1.1", "champs": {COMMENT: "Matérialité revue en 2025.\n- Pièces jointes : analyse"}},
        {"req_id": "PSG2.1-2.1.2", "champs": {COMMENT: "Revue annuelle par le comité de direction.\n- Pièces jointes : CR du 12/03/2026"}},
        {"req_id": "FR1.1-1.1.1", "champs": {COMMENT: "Société par actions simplifiée immatriculée.\n- Pièces jointes : Kbis"}},
    ], out)
    return out


def plan_rows(path):
    ws = openpyxl.load_workbook(path)[pp.SHEET]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


def test_nothing_published_without_validation(gap, tmp_path):
    rows = plan_rows(pp.build(gap, tmp_path / "plan.xlsx"))
    assert rows == []


def test_only_validated_comments_grouped_by_platform_code(gap, tmp_path):
    set_status(gap, "Propositions", lambda v: v["req_id"] in ("PSG1.1-1.1.1", "PSG2.1-2.1.1", "PSG2.1-2.1.2"), "Validé")
    rows = plan_rows(pp.build(gap, tmp_path / "plan.xlsx"))
    assert [r["Code plateforme"] for r in rows] == ["PSG1.1", "PSG2.1"]
    assert "Critère" not in rows[0]["Commentaire à coller"]          # un seul critère : texte brut
    text = rows[1]["Commentaire à coller"]
    assert text.index("Critère 2.1.1") < text.index("Critère 2.1.2")
    assert "Revue annuelle" in text and "Matérialité" in text
    assert rows[0]["Lot"] == "PSG" and rows[0]["Statut cible"] == "En cours"


def test_rejected_comment_is_excluded(gap, tmp_path):
    set_status(gap, "Propositions", lambda v: v["req_id"] == "PSG2.1-2.1.1", "Validé")
    set_status(gap, "Propositions", lambda v: v["req_id"] == "PSG2.1-2.1.2", "Rejeté")
    text = plan_rows(pp.build(gap, tmp_path / "plan.xlsx"))[0]["Commentaire à coller"]
    assert "Revue annuelle" not in text


def test_human_written_comment_counts_as_validated(tmp_path):
    src = bg.build(PROFILE, tmp_path / "gap.xlsx")
    wb = openpyxl.load_workbook(src)
    ws = wb["3. Gap analysis"]
    headers = [c.value for c in ws[1]]
    r = next(i for i in range(2, ws.max_row + 1) if ws.cell(i, headers.index("req_id") + 1).value == "FR1.1-1.1.1")
    ws.cell(r, headers.index(COMMENT) + 1).value = "Texte saisi par la consultante."
    wb.save(src)
    rows = plan_rows(pp.build(src, tmp_path / "plan.xlsx"))
    assert rows[0]["Code plateforme"] == "FR1.1"


def test_lot_filter(gap, tmp_path):
    set_status(gap, "Propositions", lambda v: True, "Validé")
    rows = plan_rows(pp.build(gap, tmp_path / "plan.xlsx", lot="FR"))
    assert {r["Lot"] for r in rows} == {"FR"}


def test_files_and_evidence_comments_from_evidence_plan(gap, tmp_path):
    set_status(gap, "Propositions", lambda v: True, "Validé")
    d = tmp_path / "client"
    d.mkdir()
    (d / "Statuts ClientTest.pdf").write_bytes(b"s")
    evid = ep.build_plan(gap, d, tmp_path / "preuves.xlsx")
    set_status(evid, "Plan de preuves", lambda v: v["Statut"] == "Trouvé", "Validé", col="Décision")
    rows = plan_rows(pp.build(gap, tmp_path / "plan.xlsx", evidence_plan=evid))
    psg = next(r for r in rows if r["Code plateforme"] == "PSG1.1")
    assert "MGPP 1.1/MGPP1.1_1.1.1_Statuts-ClientTest.pdf" in psg["Fichiers à attacher"]


def test_refuses_overwrite(gap, tmp_path):
    out = tmp_path / "plan.xlsx"
    out.write_bytes(b"x")
    with pytest.raises(FileExistsError):
        pp.build(gap, out)
