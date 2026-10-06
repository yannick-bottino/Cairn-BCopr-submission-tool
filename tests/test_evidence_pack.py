import csv
import json
import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills" / "bcorp-gap-analysis" / "scripts"))
sys.path.insert(0, str(ROOT / "skills" / "bcorp-evidence-pack" / "scripts"))

import build_gap_excel as bg  # noqa: E402
import fill_gap as fg  # noqa: E402
import evidence_plan as ep  # noqa: E402
import apply_reorg as ar  # noqa: E402
from test_build_gap_excel import PROFILE  # noqa: E402

PREUVES = "Preuves (& intitulés associés)"


@pytest.fixture
def dossier(tmp_path):
    d = tmp_path / "preuves_client"
    (d / "Juridique").mkdir(parents=True)
    (d / "RH").mkdir()
    (d / "Juridique" / "Statuts MAJ au 15.01.2024 (ClientTest) vdef.pdf").write_bytes(b"statuts")
    (d / "Juridique" / "Kbis 2026.pdf").write_bytes(b"kbis")
    (d / "RH" / "Charte télétravail v1.docx").write_bytes(b"v1")
    (d / "RH" / "Charte télétravail v2.docx").write_bytes(b"v2")
    (d / ".DS_Store").write_bytes(b"x")
    return d


@pytest.fixture
def gap(tmp_path):
    src = bg.build(PROFILE, tmp_path / "gap.xlsx")
    props = [
        {"req_id": "PSG1.1-1.1.1", "champs": {PREUVES: "Preuves attendues :\n- Statuts ClientTest MAJ 15/01/2024\n- Page raison d'être du site avec capture datée"}},
        {"req_id": "FR1.1-1.1.1", "champs": {PREUVES: "Preuves attendues :\n- Extrait Kbis de moins de 3 mois"}},
        {"req_id": "FW1.1-1.1.1", "champs": {PREUVES: "- Charte télétravail"}},
    ]
    out = tmp_path / "gap_rempli.xlsx"
    fg.fill(src, props, out)
    return out


def plan_rows(path):
    ws = openpyxl.load_workbook(path)["Plan de preuves"]
    headers = [c.value for c in ws[1]]
    return [dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]


@pytest.mark.parametrize("cell, items", [
    ("Preuves attendues :\n- A\n- B", ["A", "B"]),
    ("cf. drive MGPP 1.1\n> Statuts\n> PV d'AG", ["Statuts", "PV d'AG"]),
    ("Statuts signés", ["Statuts signés"]),
    ("", []),
])
def test_parse_evidence_cell(cell, items):
    assert ep.parse_items(cell) == items


def test_plan_matches_files(gap, dossier, tmp_path):
    out = ep.build_plan(gap, dossier, tmp_path / "plan.xlsx")
    rows = plan_rows(out)
    by = {(r["req_id"], r["Preuve attendue"]): r for r in rows}
    statuts = by[("PSG1.1-1.1.1", "Statuts ClientTest MAJ 15/01/2024")]
    assert statuts["Statut"] == "Trouvé"
    assert statuts["Fichier source"].endswith("Statuts MAJ au 15.01.2024 (ClientTest) vdef.pdf")
    assert by[("PSG1.1-1.1.1", "Page raison d'être du site avec capture datée")]["Statut"] == "Manquant"
    assert by[("FR1.1-1.1.1", "Extrait Kbis de moins de 3 mois")]["Statut"] == "Trouvé"
    assert by[("FW1.1-1.1.1", "Charte télétravail")]["Statut"] == "Ambigu"   # v1 et v2
    assert all(r["Décision"] == "À valider" for r in rows)
    assert not any(".DS_Store" in (r["Fichier source"] or "") for r in rows)


def test_proposed_destination(gap, dossier, tmp_path):
    rows = plan_rows(ep.build_plan(gap, dossier, tmp_path / "plan.xlsx"))
    statuts = next(r for r in rows if r["Statut"] == "Trouvé" and r["Code plateforme"] == "PSG1.1")
    dest = statuts["Chemin proposé"]
    assert dest.startswith("1. MGPP Mission et gouvernance des parties prenantes/MGPP 1.1/")
    assert dest.endswith(".pdf") and "MGPP1.1_1.1.1_" in dest


def test_evidence_comments_are_attached(gap, dossier, tmp_path):
    comments = tmp_path / "c.json"
    comments.write_text(json.dumps([{"req_id": "PSG1.1-1.1.1", "preuve": "Statuts ClientTest MAJ 15/01/2024",
                                     "commentaire": "Statuts à jour, raison d'être art. 2 p.1."}], ensure_ascii=False),
                        encoding="utf-8")
    rows = plan_rows(ep.build_plan(gap, dossier, tmp_path / "plan.xlsx", comments))
    r = next(r for r in rows if r["Preuve attendue"].startswith("Statuts"))
    assert r["Commentaire de preuve"] == "Statuts à jour, raison d'être art. 2 p.1."


def test_comment_style_is_checked(gap, dossier, tmp_path):
    comments = tmp_path / "c.json"
    comments.write_text(json.dumps([{"req_id": "PSG1.1-1.1.1", "preuve": "Statuts ClientTest MAJ 15/01/2024",
                                     "commentaire": "Voir CA4.4."}]), encoding="utf-8")
    with pytest.raises(ValueError, match="CA4.4"):
        ep.build_plan(gap, dossier, tmp_path / "plan.xlsx", comments)


def test_source_folder_is_read_only(gap, dossier, tmp_path):
    before = sorted(p.relative_to(dossier) for p in dossier.rglob("*"))
    ep.build_plan(gap, dossier, tmp_path / "plan.xlsx")
    assert sorted(p.relative_to(dossier) for p in dossier.rglob("*")) == before
    with pytest.raises(ValueError, match="dossier source"):
        ep.build_plan(gap, dossier, dossier / "plan.xlsx")


# --- application du plan --------------------------------------------------------

def validate_all(plan, statut="Trouvé"):
    wb = openpyxl.load_workbook(plan)
    ws = wb["Plan de preuves"]
    headers = [c.value for c in ws[1]]
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, headers.index("Statut") + 1).value == statut:
            ws.cell(r, headers.index("Décision") + 1).value = "Validé"
    wb.save(plan)


def test_dry_run_copies_nothing(gap, dossier, tmp_path):
    plan = ep.build_plan(gap, dossier, tmp_path / "plan.xlsx")
    validate_all(plan)
    dest = tmp_path / "2. Preuves (par thématique)"
    actions = ar.apply(plan, dossier, dest, execute=False)
    assert len(actions) == 2 and not dest.exists()


def test_execute_copies_validated_only_and_logs(gap, dossier, tmp_path):
    plan = ep.build_plan(gap, dossier, tmp_path / "plan.xlsx")
    validate_all(plan)
    dest = tmp_path / "2. Preuves (par thématique)"
    ar.apply(plan, dossier, dest, execute=True)
    copied = sorted(p.name for p in dest.rglob("*") if p.is_file() and p.name != ar.JOURNAL)
    assert len(copied) == 2
    assert (dossier / "Juridique" / "Kbis 2026.pdf").exists()          # copie, jamais déplacement
    journal = list(csv.DictReader(open(dest / ar.JOURNAL, encoding="utf-8")))
    assert len(journal) == 2 and all(len(j["sha256"]) == 64 for j in journal)


def test_execute_never_overwrites(gap, dossier, tmp_path):
    plan = ep.build_plan(gap, dossier, tmp_path / "plan.xlsx")
    validate_all(plan)
    dest = tmp_path / "out"
    ar.apply(plan, dossier, dest, execute=True)
    second = ar.apply(plan, dossier, dest, execute=True)
    assert all(a.result == "déjà présent" for a in second)


def test_destination_inside_source_refused(gap, dossier, tmp_path):
    plan = ep.build_plan(gap, dossier, tmp_path / "plan.xlsx")
    with pytest.raises(ValueError, match="dossier source"):
        ar.apply(plan, dossier, dossier / "rangé", execute=True)
