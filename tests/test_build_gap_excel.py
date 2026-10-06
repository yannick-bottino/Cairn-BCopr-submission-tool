import sys
from pathlib import Path

import openpyxl
import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "skills" / "bcorp-gap-analysis" / "scripts"))

import bcorp_ref as ref  # noqa: E402
import build_gap_excel as bg  # noqa: E402

PROFILE = {
    "client": "ClientTest",
    "taille": "Medium",
    "secteur": "Wholesale/Retail",
    "industrie": "",
    "horizon": 3,
    "mecanisme_equite": "À confirmer",
    "date_depot": "Avril 2027",
    "co_prestataire": "Atelier X",
    "options_retenues": ["JEDI2.g"],
    "source_profil": "Export PDF plateforme B Lab",
}


@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    out = tmp_path_factory.mktemp("o") / "gap.xlsx"
    bg.build(PROFILE, out)
    return openpyxl.load_workbook(out)


@pytest.fixture(scope="module")
def gap(wb):
    ws = wb["3. Gap analysis"]
    headers = [c.value for c in ws[1]]
    rows = [dict(zip(headers, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    return ws, headers, rows


def test_sheets(wb):
    assert wb.sheetnames == [
        "0. Mode d'emploi", "1. Paramètres client", "2. Référentiel B Corp V2.2",
        "3. Gap analysis", "4. Couverture feuille de route", "5. Répartition des rôles",
        "6. Récap gap analysis", "_meta",
    ]
    assert wb["_meta"].sheet_state == "hidden"


def test_gap_headers(gap):
    _, headers, _ = gap
    assert headers == [
        "Impact Area", "Exigence (thématique)", "Code exigence", "Code plateforme",
        "Sous-exigence", "Critère de conformité", "Clarification / informations complémentaires",
        "Année", "Niveau de conformité", "Diagnostic & gap analysis", "Actions recommandées",
        "n° action feuille de route ClientTest", "Priorité feuille de route (client)",
        "Priorité B Corp (Anchor)", "Deadline", "Statut", "Responsable", "Équipe projet",
        "Preuves (& intitulés associés)", "Commentaire soumission dossier pour l'auditeur",
        "Plateforme : commentaire ajouté", "Plateforme : statut", "Plateforme : preuves ajoutées",
        "Commentaires", "Typologie du gap (Anchor)", "Criticité (Anchor)",
        "Justification de la typologie (Anchor)", "Livrable(s) Atelier X mobilisé(s)",
        "Couverture par les livrables Atelier X",
        "Condition ou reste à produire (hors livrables Atelier X)", "req_id",
    ]


def test_every_code_comes_from_referentiel(gap):
    r = ref.Referentiel.load()
    _, _, rows = gap
    for row in rows:
        assert r.to_site(row["Code exigence"]) == row["Code plateforme"]
        assert r.to_excel(row["Code plateforme"]) == row["Code exigence"]


def test_req_ids_unique_and_complete(gap):
    r = ref.Referentiel.load()
    _, _, rows = gap
    ids = [row["req_id"] for row in rows]
    assert len(ids) == len(set(ids))
    expected = [i for x in r.applicable("Medium", "Wholesale/Retail", 3)
                for i in ([x.code] if x.type_ligne == "question_risk_tool"
                          else [f"{x.code}-{c.id}" for c in x.criteria_until(3)])]
    assert sorted(ids) == sorted(expected)


def test_sub_requirement_counts_for_medium_wr_y3(gap):
    _, _, rows = gap
    r = ref.Referentiel.load()
    subs = {row["Code plateforme"] for row in rows if r.get(row["Code plateforme"]).type_ligne == "sous_exigence"}
    years = [r.get(c).year for c in subs]
    assert (years.count(0), years.count(3), years.count(5)) == (36, 18, 0)


def test_both_codes_and_years_are_filled(gap):
    _, _, rows = gap
    assert all(row["Code exigence"] and row["Code plateforme"] for row in rows)
    assert {row["Année"] for row in rows} == {"Before Year 0", "Year 0", "Year 3"}


def test_unretained_options_are_hidden(gap):
    ws, headers, rows = gap
    for i, row in enumerate(rows, start=2):
        code = row["Code plateforme"]
        if code.startswith("JEDI2.") and code[-1].isalpha():
            assert ws.row_dimensions[i].hidden is (code != "JEDI2.g"), code


def test_drop_down_lists(gap):
    ws, headers, _ = gap
    col = {h: openpyxl.utils.get_column_letter(i + 1) for i, h in enumerate(headers)}
    lists = {}
    for dv in ws.data_validations.dataValidation:
        for rng in str(dv.sqref).split():
            lists[rng.split(":")[0].rstrip("0123456789")] = dv.formula1
    assert lists[col["Niveau de conformité"]] == '"0,1,2,NA"'
    assert lists[col["Priorité B Corp (Anchor)"]] == '"Critique,Haute,Moyenne,Basse,NA"'
    assert lists[col["Plateforme : statut"]] == '"Non démarré,En cours,Terminé : Preuve Prêt"'
    assert "1-Critique" in lists[col["Criticité (Anchor)"]]


def test_writer_columns_are_empty(gap):
    _, _, rows = gap
    for h in ("Niveau de conformité", "Diagnostic & gap analysis", "Actions recommandées",
              "Commentaire soumission dossier pour l'auditeur"):
        assert all(row[h] in (None, "") for row in rows)


def test_parameters_sheet(wb):
    ws = wb["1. Paramètres client"]
    params = {r[0].value: r[1].value for r in ws.iter_rows(min_row=2) if r[0].value}
    assert params["Client"] == "ClientTest"
    assert params["Taille (B Lab)"] == "Medium"
    assert params["Horizon retenu"] == "Year 0 + Year 3"


def test_recap_uses_formulas(wb):
    ws = wb["6. Récap gap analysis"]
    formulas = [c.value for row in ws.iter_rows() for c in row if isinstance(c.value, str) and c.value.startswith("=")]
    counts = [f for f in formulas if f.startswith("=COUNT")]
    assert counts and all("'3. Gap analysis'" in f for f in counts)


def test_refuses_to_overwrite(tmp_path):
    out = tmp_path / "gap.xlsx"
    out.write_bytes(b"existing")
    with pytest.raises(FileExistsError):
        bg.build(PROFILE, out)


def test_rejects_unknown_option(tmp_path):
    with pytest.raises(ref.UnknownCode):
        bg.build({**PROFILE, "options_retenues": ["JEDI2.zz"]}, tmp_path / "x.xlsx")


def by_req_id(rows):
    return {row["req_id"]: row for row in rows}


def test_criterion_deadline_overrides_sub_requirement_year(gap):
    _, _, rows = gap
    ids = by_req_id(rows)
    assert ids["ESC1.4-1.4.2"]["Année"] == "Before Year 0"
    assert "ESC1.1-1.1.7" not in ids          # « For Year 5 » : hors horizon Year 3


def test_year5_criteria_included_at_horizon_5(tmp_path):
    out = bg.build({**PROFILE, "horizon": 5}, tmp_path / "g5.xlsx")
    ws = openpyxl.load_workbook(out)["3. Gap analysis"]
    headers = [c.value for c in ws[1]]
    ids = {dict(zip(headers, [c.value for c in r]))["req_id"]: r for r in ws.iter_rows(min_row=2)}
    assert "ESC1.1-1.1.7" in ids


def test_criterion_text_is_filled_fr_first(gap):
    _, _, rows = gap
    ids = by_req_id(rows)
    psg = ids["PSG1.1-1.1.1"]["Critère de conformité"]
    assert psg.startswith("1.1.1") and "raison d'être" in psg
    en_only = [row["Critère de conformité"] for row in rows
               if row["Critère de conformité"] and "[EN]" in row["Critère de conformité"]]
    assert en_only, "les critères sans traduction gardent le texte EN, signalé [EN]"
    assert not any("texte à reprendre du PDF" in (row["Critère de conformité"] or "") for row in rows)


def test_pdf_typo_is_corrected(gap):
    _, _, rows = gap
    ids = by_req_id(rows)
    assert "FW1.1-1.1.3" in ids and "FW1.1-1.2.3" not in ids


def test_profile_source_is_declared_not_computed(wb):
    ws = wb["1. Paramètres client"]
    params = {r[0].value: r[1].value for r in ws.iter_rows(min_row=2) if r[0].value}
    assert params["Source du profil (taille, secteur)"] == "Export PDF plateforme B Lab"
    assert "Effectif" not in params and "Chiffre d'affaires" not in params


def test_profile_source_is_required(tmp_path):
    profile = {k: v for k, v in PROFILE.items() if k != "source_profil"}
    with pytest.raises(ValueError, match="source_profil"):
        bg.build(profile, tmp_path / "x.xlsx")


def test_profile_source_values(tmp_path):
    with pytest.raises(ValueError, match="source_profil"):
        bg.build({**PROFILE, "source_profil": "calculé"}, tmp_path / "x.xlsx")
