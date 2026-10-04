#!/usr/bin/env python3
"""Génère l'Excel de gap analysis B Corp (format gabarit Anchor à 8 onglets, codes FR + EN).

Usage :
    python3 build_gap_excel.py <profil.json> <sortie.xlsx>

Le script ne remplit que le référentiel filtré (taille, secteur, horizon) et la structure.
Les colonnes rédactionnelles (niveau, diagnostic, actions, commentaire auditeur) restent vides :
elles sont proposées par Claude puis validées par la consultante. Il refuse d'écraser un fichier.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parents[3] / "shared" / "lib"))
import bcorp_ref as ref  # noqa: E402

SCHEMA_VERSION = "0.1.0"

HEADER_FILL = PatternFill("solid", fgColor="1F3B5C")
PLATFORM_FILL = PatternFill("solid", fgColor="7F7F7F")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=10)
BODY_FONT = Font(name="Calibri", size=9)
WRAP = Alignment(wrap_text=True, vertical="top")

LISTS = {
    "Niveau de conformité": ["0", "1", "2", "NA"],
    "Priorité feuille de route (client)": ["1", "2", "3", "NA"],
    "Priorité B Corp (Anchor)": ["Critique", "Haute", "Moyenne", "Basse", "NA"],
    "Statut": ["À lancer", "En cours", "Terminé", "NA"],
    "Plateforme : commentaire ajouté": ["Oui", "Non"],
    # Statuts réellement affichés par app.bcorporation.net (testés les 17-18/09/2026)
    "Plateforme : statut": ["Non démarré", "En cours", "Terminé : Preuve Prêt"],
    "Plateforme : preuves ajoutées": ["Oui", "Partiel", "Non"],
    "Typologie du gap (Anchor)": [
        "Correctif mineur / pièce à produire", "Process à mettre en place",
        "Document structurant", "Chantier majeur - arbitrage CODIR", "NA",
    ],
    "Criticité (Anchor)": ["1-Critique", "2-Majeur", "3-Modéré", "4-Mineur", "NA"],
}
CRITICITY_COLORS = {"1-Critique": "C0392B", "2-Majeur": "E67E22", "3-Modéré": "F4D03F",
                    "4-Mineur": "7FB77E", "NA": "D9D9D9"}
WIDTHS = {"Impact Area": 22, "Exigence (thématique)": 30, "Code exigence": 11, "Code plateforme": 11,
          "Sous-exigence": 40, "Critère de conformité": 56,
          "Clarification / informations complémentaires": 40, "Diagnostic & gap analysis": 69,
          "Actions recommandées": 63, "Preuves (& intitulés associés)": 45,
          "Commentaire soumission dossier pour l'auditeur": 92, "Commentaires": 40,
          "Justification de la typologie (Anchor)": 40}
HORIZON_LABELS = {0: "Year 0", 3: "Year 0 + Year 3", 5: "Year 0 + Year 3 + Year 5"}


def gap_headers(client: str, co: str) -> list[str]:
    return [
        "Impact Area", "Exigence (thématique)", "Code exigence", "Code plateforme",
        "Sous-exigence", "Critère de conformité", "Clarification / informations complémentaires",
        "Année", "Niveau de conformité", "Diagnostic & gap analysis", "Actions recommandées",
        f"n° action feuille de route {client}", "Priorité feuille de route (client)",
        "Priorité B Corp (Anchor)", "Deadline", "Statut", "Responsable", "Équipe projet",
        "Preuves (& intitulés associés)", "Commentaire soumission dossier pour l'auditeur",
        "Plateforme : commentaire ajouté", "Plateforme : statut", "Plateforme : preuves ajoutées",
        "Commentaires", "Typologie du gap (Anchor)", "Criticité (Anchor)",
        "Justification de la typologie (Anchor)", f"Livrable(s) {co} mobilisé(s)",
        f"Couverture par les livrables {co}",
        f"Condition ou reste à produire (hors livrables {co})", "req_id",
    ]


def build(profile: dict, out: Path) -> Path:
    out = Path(out)
    if out.exists():
        raise FileExistsError(f"{out} existe déjà : le script n'écrase jamais un fichier.")
    r = ref.Referentiel.load()
    size, sector, horizon = profile["taille"], profile["secteur"], int(profile["horizon"])
    retained = {r.to_site(c) for c in profile.get("options_retenues", [])}
    rows = r.applicable(size, sector, horizon)
    rows.sort(key=lambda x: r.order.index(x.impact_area))  # tri stable : garde l'ordre du PDF

    wb = Workbook()
    wb.remove(wb.active)
    _mode_emploi(wb.create_sheet("0. Mode d'emploi"))
    _parametres(wb.create_sheet("1. Paramètres client"), profile)
    _referentiel(wb.create_sheet("2. Référentiel B Corp V2.2"), r, size, sector, horizon)
    headers = gap_headers(profile["client"], profile.get("co_prestataire") or "co-prestataire")
    _gap(wb.create_sheet("3. Gap analysis"), r, rows, retained, headers)
    _fdr(wb.create_sheet("4. Couverture feuille de route"))
    _roles(wb.create_sheet("5. Répartition des rôles"), r, profile)
    _recap(wb.create_sheet("6. Récap gap analysis"), r, headers)
    _meta(wb.create_sheet("_meta"), profile)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


# --- onglets -----------------------------------------------------------------

def _header(ws, headers, widths=None, platform=()):
    ws.append(headers)
    for i, h in enumerate(headers, start=1):
        c = ws.cell(row=1, column=i)
        c.font, c.alignment = HEADER_FONT, WRAP
        c.fill = PLATFORM_FILL if h in platform else HEADER_FILL
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(h, 16)
    ws.freeze_panes = "A2"


def _gap(ws, r, rows, retained, headers):
    platform = [h for h in headers if h.startswith("Plateforme :")]
    _header(ws, headers, WIDTHS, platform)
    ws.freeze_panes = "F2"
    col = {h: i for i, h in enumerate(headers, start=1)}
    tone, last_req = {}, None
    for x in rows:
        ia = r.impact_area(x.impact_area)
        if x.requirement_code != last_req:
            tone[x.impact_area] = 1 - tone.get(x.impact_area, 1)
            last_req = x.requirement_code
        fill = PatternFill("solid", fgColor=ia.colors[tone[x.impact_area] % len(ia.colors)])
        n = x.requirement_code[len(ia.code):]
        if x.type_ligne == "question_risk_tool":
            # Question oui/non du Risk Tool : une seule ligne, les exigences déclenchées sont listées à part
            units = [(f"Question Risk Tool (oui/non sur la plateforme). Critères : {', '.join(x.criteria_ids)}", x.code)]
        else:
            units = [(f"{cid} [texte à reprendre du PDF V2.2, p.{x.page_pdf}]", rid)
                     for cid, rid in zip(x.criteria_ids or [""], x.req_ids)]
        for criterion, rid in units:
            values = {
                "Impact Area": f"{ia.name_fr} ({ia.prefix_excel})",
                "Exigence (thématique)": f"{ia.prefix_excel} {n} : {x.requirement_en}",
                "Code exigence": x.code_excel,
                "Code plateforme": x.code,
                "Sous-exigence": x.title_fr or f"[EN] {x.title_en}",
                "Critère de conformité": criterion,
                "Année": f"Year {x.year}",
                "req_id": rid,
            }
            ws.append([values.get(h) for h in headers])
            i = ws.max_row
            for c in ws[i]:
                c.font, c.alignment, c.fill = BODY_FONT, WRAP, fill
            if x.type_ligne == "option" and x.code not in retained:
                ws.row_dimensions[i].hidden = True
    last = max(ws.max_row, 2)
    for h, values in LISTS.items():
        letter = get_column_letter(col[h])
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
        dv.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv)
    crit = get_column_letter(col["Criticité (Anchor)"])
    for value, color in CRITICITY_COLORS.items():
        ws.conditional_formatting.add(
            f"{crit}2:{crit}{last}",
            CellIsRule(operator="equal", formula=[f'"{value}"'], fill=PatternFill("solid", fgColor=color)))
    ws.column_dimensions[get_column_letter(col["req_id"])].hidden = True
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{last}"


def _mode_emploi(ws):
    lines = [
        "Gap analysis B Corp, B Lab Standards V2.2 (20/02/2026). Anchor Strategy B Corp tool.",
        "",
        "1. L'onglet 1 fixe le périmètre (taille, secteur, horizon). Pour le changer, régénérer le fichier.",
        "2. L'onglet 3 liste un critère de conformité par ligne. Code exigence = sigle FR, Code plateforme = code affiché sur app.bcorporation.net.",
        "3. Les options de menu non retenues (JEDI 2, APAC 2…) sont masquées, pas supprimées.",
        "4. Niveau de conformité : 0 = non couvert, 1 = partiellement couvert, 2 = pleinement couvert, NA = non applicable.",
        "5. Les colonnes Diagnostic, Actions, Priorité, Responsable et Commentaires ne sont jamais publiées. Seul le commentaire auditeur l'est, après validation.",
        "6. La colonne req_id (masquée) est la clé technique entre modules : ne pas la modifier.",
    ]
    for line in lines:
        ws.append([line])
    ws.column_dimensions["A"].width = 140


def _parametres(ws, p):
    _header(ws, ["Paramètre", "Valeur", "Liste / règle"], {"Paramètre": 32, "Valeur": 30, "Liste / règle": 90})
    rows = [
        ("Client", p["client"], "texte libre"),
        ("Taille (B Lab)", p["taille"], ", ".join(ref.SIZES) + ". Règle : la plus petite des deux tailles (effectif, CA)."),
        ("Statut de la taille", p.get("statut_taille", "À confirmer"), "Confirmée seulement si les seuils sont sourcés."),
        ("Effectif", p.get("effectif", ""), "par pays si possible"),
        ("Chiffre d'affaires", p.get("ca", ""), "dernier exercice clos, devise précisée"),
        ("Secteur (B Lab)", p["secteur"], ", ".join(ref.SECTORS)),
        ("Industrie", p.get("industrie", ""), "seulement si l'industrie déclenche des exigences spécifiques"),
        ("Horizon retenu", HORIZON_LABELS[int(p["horizon"])], ", ".join(HORIZON_LABELS.values())),
        ("Mécanisme d'équité applicable", p.get("mecanisme_equite", "À confirmer"), "Oui, Non, À confirmer"),
        ("Date de dépôt visée", p.get("date_depot", ""), "texte libre"),
        ("Co-prestataire", p.get("co_prestataire", ""), "producteur de preuves tiers, s'il existe"),
        ("Options retenues", ", ".join(p.get("options_retenues", [])), "codes plateforme, ex. JEDI2.g"),
    ]
    for row in rows:
        ws.append(list(row))
        for c in ws[ws.max_row]:
            c.font, c.alignment = BODY_FONT, WRAP


def _referentiel(ws, r, size, sector, horizon):
    headers = ["Applicable ?", "Impact Area (FR)", "Topic (EN)", "Code exigence", "Code plateforme", "Type",
               "Intitulé sous-exigence (FR)", "Sub-requirement title (EN)", "Critères de conformité",
               "Année", "Page PDF", "Exemples de preuves (KB, à relire)", "Statut vérification"]
    _header(ws, headers, {"Intitulé sous-exigence (FR)": 50, "Sub-requirement title (EN)": 50,
                          "Exemples de preuves (KB, à relire)": 50})
    ws.freeze_panes = "E2"
    applicable = {x.code for x in r.applicable(size, sector, horizon)}
    for x in r.rows:
        ia = r.impact_area(x.impact_area)
        ws.append(["OUI" if x.code in applicable else "non", ia.name_fr, ia.name_en, x.code_excel, x.code,
                   x.type_ligne, x.title_fr, x.title_en, ", ".join(x.criteria_ids), f"Year {x.year}",
                   x.page_pdf, x.evidence_examples, "à vérifier / PDF officiel"])
        for c in ws[ws.max_row]:
            c.font = BODY_FONT
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{ws.max_row}"


def _fdr(ws):
    headers = ["n° action FDR", "Intitulé de l'action (texte de la FDR)", "Code(s) plateforme servi(s)",
               "Lecture : couverture", "Justification (lecture du texte de l'action)"]
    _header(ws, headers, {"Intitulé de l'action (texte de la FDR)": 60,
                          "Justification (lecture du texte de l'action)": 60})
    dv = DataValidation(type="list", formula1='"Couvert,Partiel,Non couvert,Formalité"', allow_blank=True)
    dv.add("D2:D500")
    ws.add_data_validation(dv)


def _roles(ws, r, p):
    co = p.get("co_prestataire") or "Co-prestataire"
    _header(ws, ["Impact Area", "Producteur de la preuve", "Zone de recouvrement à trancher", "Décision (qui, quand)"],
            {"Impact Area": 45, "Zone de recouvrement à trancher": 50, "Décision (qui, quand)": 40})
    for code in r.order:
        ia = r.impact_area(code)
        ws.append([f"{ia.name_fr} ({ia.prefix_excel})"])
    dv = DataValidation(type="list", formula1=f'"Anchor Strategy,{p["client"]},{co},À trancher"', allow_blank=True)
    dv.add(f"B2:B{ws.max_row}")
    ws.add_data_validation(dv)


def _recap(ws, r, headers):
    col = {h: get_column_letter(i) for i, h in enumerate(headers, start=1)}
    g = "'3. Gap analysis'"
    ia_rng = f"{g}!${col['Impact Area']}:${col['Impact Area']}"
    lvl_rng = f"{g}!${col['Niveau de conformité']}:${col['Niveau de conformité']}"
    crit_rng = f"{g}!${col['Criticité (Anchor)']}:${col['Criticité (Anchor)']}"
    _header(ws, ["Impact Area", "Critères", "0 = non couvert", "1 = partiel", "2 = couvert", "NA",
                 "Non renseigné", "1-Critique", "2-Majeur"], {"Impact Area": 50})
    for code in r.order:
        ia = r.impact_area(code)
        label = f"{ia.name_fr} ({ia.prefix_excel})"
        n = ws.max_row + 1
        ws.append([
            label,
            f'=COUNTIF({ia_rng},$A{n})',
            f'=COUNTIFS({ia_rng},$A{n},{lvl_rng},"0")+COUNTIFS({ia_rng},$A{n},{lvl_rng},0)',
            f'=COUNTIFS({ia_rng},$A{n},{lvl_rng},"1")+COUNTIFS({ia_rng},$A{n},{lvl_rng},1)',
            f'=COUNTIFS({ia_rng},$A{n},{lvl_rng},"2")+COUNTIFS({ia_rng},$A{n},{lvl_rng},2)',
            f'=COUNTIFS({ia_rng},$A{n},{lvl_rng},"NA")',
            f'=COUNTIFS({ia_rng},$A{n},{lvl_rng},"")',
            f'=COUNTIFS({ia_rng},$A{n},{crit_rng},"1-Critique")',
            f'=COUNTIFS({ia_rng},$A{n},{crit_rng},"2-Majeur")',
        ])
    n = ws.max_row
    ws.append(["Total"] + [f"=SUM({get_column_letter(c)}2:{get_column_letter(c)}{n})" for c in range(2, 10)])
    for c in ws[ws.max_row]:
        c.font = Font(name="Calibri", bold=True, size=10)


def _meta(ws, p):
    digest = hashlib.sha256(ref.CSV_PATH.read_bytes()).hexdigest()
    for k, v in [("schema_version", SCHEMA_VERSION), ("referentiel", "B Lab Standards V2.2 (20/02/2026)"),
                 ("referentiel_sha256", digest), ("generated_at", dt.datetime.now().isoformat(timespec="seconds")),
                 ("profil", json.dumps(p, ensure_ascii=False))]:
        ws.append([k, v])
    ws.sheet_state = "hidden"


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    profile = json.loads(Path(sys.argv[1]).read_text(encoding="utf-8"))
    print(build(profile, Path(sys.argv[2])))
