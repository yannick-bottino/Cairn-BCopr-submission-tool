#!/usr/bin/env python3
"""Plan de saisie app.bcorporation.net : commentaires à coller et pièces à attacher, par fiche.

Usage :
    python3 platform_plan.py <gap.xlsx> <sortie.xlsx> [--lot PSG] [--preuves plan_preuves.xlsx]

Aucune automatisation du navigateur : le plan se lit et se colle à la main, lot par lot.
Un commentaire n'entre dans le plan que s'il est validé :
- proposé par Claude (tracé dans l'onglet « Propositions ») : statut « Validé » obligatoire ;
- saisi par la consultante (absent de « Propositions ») : considéré comme validé.
Les pièces viennent du plan de preuves (lignes « Validé » seulement). Statut cible : toujours « En cours ».
"""
from __future__ import annotations

import argparse
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "bcorp-gap-analysis" / "scripts"))
import build_gap_excel as bg  # noqa: E402
import fill_gap as fg  # noqa: E402

ref = bg.ref
SHEET = "Plan de saisie"
HEADERS = ["Lot", "Code plateforme", "Code exigence", "Commentaire à coller", "Fichiers à attacher",
           "Commentaires de preuve", "Statut cible", "Fait"]
TARGET_STATUS = "En cours"   # jamais « Terminé : Preuve Prêt » : c'est le client qui bascule


def _rows(ws):
    headers = [c.value for c in ws[1]]
    for row in ws.iter_rows(min_row=2, values_only=True):
        yield dict(zip(headers, row))


def build(gap: Path, out: Path, lot: str | None = None, evidence_plan: Path | None = None) -> Path:
    out = Path(out)
    if out.exists():
        raise FileExistsError(f"{out} existe déjà : le script n'écrase jamais un fichier.")
    r = ref.Referentiel.load()
    wb = openpyxl.load_workbook(gap, data_only=True)
    status = {}
    if fg.LOG_SHEET in wb.sheetnames:
        for p in _rows(wb[fg.LOG_SHEET]):
            if p["Colonne"] == fg.COMMENT:
                status[p["req_id"]] = p["Statut"]

    fiches: "OrderedDict[str, list[tuple[str, str]]]" = OrderedDict()
    for row in _rows(wb[fg.GAP_SHEET]):
        text = (row.get(fg.COMMENT) or "").strip()
        rid = row["req_id"]
        if not text or status.get(rid, "Validé") != "Validé":
            continue
        code = row["Code plateforme"]
        if lot and r.get(code).impact_area != lot:
            continue
        fiches.setdefault(code, []).append((rid.partition("-")[2], text))

    files, evid_comments = {}, {}
    if evidence_plan:
        for e in _rows(openpyxl.load_workbook(evidence_plan, data_only=True)["Plan de preuves"]):
            if e["Décision"] == "Validé" and e["Chemin proposé"]:
                files.setdefault(e["Code plateforme"], []).append(e["Chemin proposé"])
                if e["Commentaire de preuve"]:
                    name = Path(e["Chemin proposé"]).name
                    evid_comments.setdefault(e["Code plateforme"], []).append(f"{name} : {e['Commentaire de preuve']}")

    out_wb = openpyxl.Workbook()
    ws = out_wb.active
    ws.title = SHEET
    bg._header(ws, HEADERS, {"Commentaire à coller": 90, "Fichiers à attacher": 60, "Commentaires de preuve": 60})
    for code, parts in fiches.items():
        if len(parts) == 1:
            comment = parts[0][1]
        else:
            comment = "\n\n".join(f"Critère {cid} :\n{text}" if cid else text for cid, text in parts)
        req = r.get(code)
        ws.append([req.impact_area, code, req.code_excel, comment, "\n".join(files.get(code, [])),
                   "\n".join(evid_comments.get(code, [])), TARGET_STATUS, "Non"])
        for c in ws[ws.max_row]:
            c.font, c.alignment = bg.BODY_FONT, bg.WRAP
    dv = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=False)
    dv.add(f"H2:H{max(ws.max_row, 2)}")
    ws.add_data_validation(dv)
    out.parent.mkdir(parents=True, exist_ok=True)
    out_wb.save(out)
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[1])
    ap.add_argument("gap")
    ap.add_argument("sortie")
    ap.add_argument("--lot", help="code d'Impact Area plateforme : FR, PSG, FW, JEDI, HR, CA, ESC, GACA")
    ap.add_argument("--preuves", help="plan de preuves validé (module 4)")
    a = ap.parse_args(argv)
    print(build(Path(a.gap), Path(a.sortie), a.lot, Path(a.preuves) if a.preuves else None))
    return 0


if __name__ == "__main__":
    sys.exit(main())
