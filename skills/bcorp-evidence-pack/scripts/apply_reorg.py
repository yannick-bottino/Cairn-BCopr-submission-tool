#!/usr/bin/env python3
"""Range les preuves validées du plan : copie (jamais de déplacement), journal, aucune écrasement.

Usage :
    python3 apply_reorg.py <plan.xlsx> <dossier_preuves_client> <dossier_cible>            # simulation
    python3 apply_reorg.py <plan.xlsx> <dossier_preuves_client> <dossier_cible> --execute  # copie

Seules les lignes « Décision = Validé » avec un fichier source sont traitées. Le dossier cible
ne peut pas être dans le dossier du client. Chaque copie est tracée dans _journal_copies.csv.
"""
from __future__ import annotations

import csv
import datetime as dt
import hashlib
import shutil
import sys
from dataclasses import dataclass
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
from evidence_plan import PLAN_SHEET, _inside  # noqa: E402

JOURNAL = "_journal_copies.csv"


@dataclass
class Action:
    req_id: str
    source: Path
    destination: Path
    result: str


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def apply(plan: Path, source: Path, target: Path, execute: bool = False) -> list[Action]:
    source, target = Path(source), Path(target)
    if _inside(target, source):
        raise ValueError("Le dossier cible ne peut pas être dans le dossier source du client (lecture seule).")
    ws = openpyxl.load_workbook(plan, data_only=True)[PLAN_SHEET]
    headers = [c.value for c in ws[1]]
    actions = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        rec = dict(zip(headers, row))
        if rec["Décision"] != "Validé" or not rec["Fichier source"] or not rec["Chemin proposé"]:
            continue
        src, dest = source / rec["Fichier source"], target / rec["Chemin proposé"]
        if not src.exists():
            result = "source introuvable"
        elif dest.exists():
            result = "déjà présent" if sha256(dest) == sha256(src) else "conflit : fichier différent déjà présent"
        else:
            result = "copié" if execute else "à copier"
        actions.append(Action(rec["req_id"], src, dest, result))

    if execute:
        new = [a for a in actions if a.result == "copié"]
        for a in new:
            a.destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(a.source, a.destination)
        if new:
            journal = target / JOURNAL
            first = not journal.exists()
            with open(journal, "a", newline="", encoding="utf-8") as f:
                w = csv.writer(f)
                if first:
                    w.writerow(["horodatage", "req_id", "source", "destination", "sha256"])
                now = dt.datetime.now().isoformat(timespec="seconds")
                for a in new:
                    w.writerow([now, a.req_id, str(a.source), str(a.destination.relative_to(target)), sha256(a.destination)])
    return actions


if __name__ == "__main__":
    args = [a for a in sys.argv[1:] if a != "--execute"]
    if len(args) != 3:
        sys.exit(__doc__)
    for a in apply(*[Path(x) for x in args], execute="--execute" in sys.argv):
        print(f"{a.result:40} {a.req_id:16} {a.destination}")
