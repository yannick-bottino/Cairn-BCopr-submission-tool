#!/usr/bin/env python3
"""Inventaire d'une data room client pour le diagnostic B Corp (lecture seule).

Usage :
    python3 inventory.py <dossier_client> <inventaire.xlsx>

Onglets produits :
- « Inventaire » : un fichier par ligne (chemin, type et Impact Areas pressentis, taille, date,
  empreinte SHA-256, alerte doublon / version antérieure) ;
- « Pièces clés » : pièces presque toujours demandées (Kbis, statuts, actionnariat, liasses…),
  présentes ou absentes d'après les noms de fichiers ;
- « Vue par pilier » : une ligne par Impact Area, colonnes « Ce qui existe (preuve) » et
  « Ce qui manque (Year 0) » laissées vides pour Claude et la consultante.
Le classement est un pré-tri par mots-clés (references/classification.json) : Claude arbitre.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import json
import re
import sys
import unicodedata
from pathlib import Path

from openpyxl import Workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "bcorp-gap-analysis" / "scripts"))
import build_gap_excel as bg  # noqa: E402

ref = bg.ref
RULES = json.loads((HERE.parent / "references" / "classification.json").read_text(encoding="utf-8"))
IGNORED = {".DS_Store", "Thumbs.db", "desktop.ini"}
COPY_RE = re.compile(r"\b(copie|copy|\(\d\))", re.I)
VERSION_RE = re.compile(r"[\s_-]*v(\d+)\b", re.I)


def norm(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode().lower()
    return " " + re.sub(r"[^a-z0-9]+", " ", text).strip() + " "


def hits(text: str, words: list[str]) -> int:
    """Nombre de mots-clés trouvés en début de mot (« statut » trouve « statuts »)."""
    return sum(1 for w in words if f" {w.strip()}" in text)


def classify(rel: Path) -> tuple[str, list[str]]:
    text = norm(str(rel.with_suffix("")))
    scores = {t: hits(text, words) for t, words in RULES["types"].items()}
    best = max(scores, key=lambda t: scores[t])
    areas = [code for code in ref.Referentiel.load().order if hits(text, RULES["impact_areas"].get(code, []))]
    return (best if scores[best] else "À classer"), areas


def _inside(child: Path, parent: Path) -> bool:
    try:
        child.resolve().relative_to(parent.resolve())
        return True
    except ValueError:
        return False


def list_files(root: Path) -> list[Path]:
    return sorted(p for p in root.rglob("*") if p.is_file() and p.name not in IGNORED
                  and not any(part.startswith(".") for part in p.relative_to(root).parts))


def alerts(root: Path, files: list[Path]) -> dict[Path, str]:
    out: dict[Path, list[str]] = {f: [] for f in files}
    by_hash: dict[str, list[Path]] = {}
    for f in files:
        by_hash.setdefault(hashlib.sha256(f.read_bytes()).hexdigest(), []).append(f)
    for group in by_hash.values():
        if len(group) > 1:
            original = min(group, key=lambda f: (bool(COPY_RE.search(f.stem)), len(f.name), str(f)))
            for f in group:
                if f != original:
                    out[f].append(f"doublon de {original.relative_to(root)}")
    by_stem: dict[tuple[Path, str], list[tuple[int, Path]]] = {}
    for f in files:
        m = VERSION_RE.search(f.stem)
        if m:
            key = (f.parent, VERSION_RE.sub("", f.stem).strip().lower() + f.suffix.lower())
            by_stem.setdefault(key, []).append((int(m.group(1)), f))
    for versions in by_stem.values():
        if len(versions) > 1:
            latest = max(versions)[1]
            for _, f in versions:
                if f != latest:
                    out[f].append(f"version antérieure de {latest.name}")
    return {f: " ; ".join(v) for f, v in out.items()}


def build(root: Path, out: Path) -> Path:
    root, out = Path(root), Path(out)
    if _inside(out, root):
        raise ValueError("L'inventaire ne s'écrit jamais dans le dossier client (lecture seule).")
    if out.exists():
        raise FileExistsError(f"{out} existe déjà : le script n'écrase jamais un fichier.")
    r = ref.Referentiel.load()
    files = list_files(root)
    warn = alerts(root, files)

    wb = Workbook()
    ws = wb.active
    ws.title = "Inventaire"
    bg._header(ws, ["Chemin", "Fichier", "Extension", "Type pressenti", "Impact Areas pressenties",
                    "Taille (Ko)", "Modifié le", "SHA-256", "Alerte"],
               {"Chemin": 60, "Fichier": 40, "Alerte": 45})
    area_count = {code: 0 for code in r.order}
    for f in files:
        rel = f.relative_to(root)
        kind, areas = classify(rel)
        for code in areas:
            area_count[code] += 1
        ws.append([str(rel), f.name, f.suffix.lower().lstrip("."), kind, ", ".join(areas),
                   round(f.stat().st_size / 1024, 1),
                   dt.datetime.fromtimestamp(f.stat().st_mtime).date().isoformat(),
                   hashlib.sha256(f.read_bytes()).hexdigest(), warn[f]])
    ws.auto_filter.ref = f"A1:I{max(ws.max_row, 2)}"

    keys = wb.create_sheet("Pièces clés")
    bg._header(keys, ["Pièce", "Statut", "Fichiers correspondants"], {"Pièce": 45, "Fichiers correspondants": 70})
    for item in RULES["pieces_cles"]:
        found = [str(f.relative_to(root)) for f in files if hits(norm(str(f.relative_to(root).with_suffix(""))), item["mots"])]
        keys.append([item["piece"], "Présent (à vérifier)" if found else "Absent", "\n".join(found)])

    pillars = wb.create_sheet("Vue par pilier")
    bg._header(pillars, ["Impact Area", "Documents pressentis", "Ce qui existe (preuve)", "Ce qui manque (Year 0)",
                         "Points à fiabiliser"],
               {"Impact Area": 45, "Ce qui existe (preuve)": 60, "Ce qui manque (Year 0)": 60, "Points à fiabiliser": 45})
    for code in r.order:
        ia = r.impact_area(code)
        pillars.append([f"{ia.name_fr} ({ia.prefix_excel} / {code})", area_count[code]])
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out)
    return out


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    print(build(Path(sys.argv[1]), Path(sys.argv[2])))
